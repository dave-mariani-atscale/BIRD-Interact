#!/usr/bin/env python3
"""Add the leaderboard-mode census columns to every "Leaderboard *" tab.

One new banded section per tab, following the Summary-tab conventions the rest of
the workbook uses: a merged colour-coded section header on row 4, column headers
on row 5, the 22 databases on rows 6-27, the pooled ALL row on 28 and the
best-single-run row on 29. Counts are literals from the census; every rate is a
FORMULA over columns already on the tab, so the sheet stays auditable and
recalculates if a number changes.

  Query P1 / P2 unwinnable Qs   census revision LB1, per database
  Query P1 unwinnable %         unwinnable / Query Qs
  Query P1 / P2 Adjusted %      passes / (Query Qs - unwinnable): the pass rate
                                over the tasks a correct answer can actually win
                                under the board's own comparison
  Success Rate / Blended P2 Adjusted   the same adjustment on the leaderboard's
                                own blended denominator (600 tasks). Only the
                                Query side is censused, so only Query unwinnables
                                come out of the denominator; Management tasks are
                                all left in.
  Raw Query P1 / P2 Adjusted %  the same for the raw control arm, where the tab
                                carries one
  Query P1 / P2 adjusted lift   AtScale adjusted rate / raw adjusted rate - 1.
                                Identical by construction to the plain Query lift
                                (both arms divide by the same winnable count, so
                                it cancels); carried so the lift can be read in
                                the same units as the rates beside it
  Board-broken golds (P1)       tasks whose gold UPSTREAM's own cleanup turns into
                                invalid SQL, so the board scores 0 whatever is
                                submitted. Not deducted anywhere - this harness
                                keeps the two DISTINCT operators intact and DOES
                                score them, which is the one place our figures sit
                                ABOVE what the board would compute.

    PYTHONPATH=. .venv-adk/bin/python scripts/add_leaderboard_census_columns.py \
        --census /Users/.../bird-atscale-models/census/census_lb1_final.json --in-place
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as GL

DRIVE_XLSX = ("/Users/davidmariani/Library/CloudStorage/GoogleDrive-dave@atscale.com/"
              "Shared drives/Product/Benchmark/BIRD Benchmark/Results/"
              "BIRD Results (atscale n=3, memory=shapes, mcp=feedback-memory-combined, engine=pr9967).xlsx")
FIRST, LAST, ALL, BEST, NOTE = 6, 27, 28, 29, 30
NAVY, LILAC, BAND = "FF1F3864", "FFD9D2E9", "FFFFF2CC"
PCT = "0.0%"

HOW_TO_READ = (
    " CENSUS-ADJUSTED (added 2026-09-16, census revision {rev} over {runs} leaderboard-mode runs, "
    "docs/LEADERBOARD_TASK_CENSUS.md). UNWINNABLE Qs are Query tasks that no correct answer can win under the "
    "leaderboard protocol - a verified gold defect, an answer that depends on the query plan, or gold output "
    "(text casing, column order, row order without an ordering cue) that upstream's exact comparison demands and "
    "nothing in the question, knowledge base or simulator discloses. Every one of them also has ZERO passes across "
    "every leaderboard-mode run recorded to date, both arms and every agent model, so the count is a floor, not a "
    "claim. ADJUSTED % divides the same passes by the winnable tasks only; it is the pass rate over the part of the "
    "benchmark that can be won, and it is NOT comparable to the board's published Success Rate, which keeps every "
    "task in the denominator. The blended adjusted figures take only Query unwinnables out of the 600-task "
    "denominator - Management tasks have no census. QUERY P1/P2 ADJUSTED LIFT divides the AtScale "
    "adjusted rate by the raw arm's; it comes out identical to the plain Query lift because the "
    "winnable-task denominator cancels in a ratio, which is the point of carrying it. BOARD-BROKEN GOLDS are the other direction: upstream's own "
    "remove_distinct deletes the word DISTINCT out of DISTINCT ON (...) and IS [NOT] DISTINCT FROM, so the BOARD "
    "executes invalid gold and scores 0 no matter what is submitted, while this harness keeps both operators and "
    "scores those tasks normally. They are not deducted from anything here; they are the one place these figures sit "
    "above what the board would compute.")


MARKER = " CENSUS-ADJUSTED (added 2026-09-16"


def upsert_note(ws, marker, text, next_markers=()):
    """Replace this script's own paragraph in the HOW TO READ cell, leaving any
    other script's paragraph alone - so the two tab maintenance scripts can run
    in either order without eating each other's note."""
    cur = ws[f"A{NOTE}"].value or ""
    i = cur.find(marker)
    if i < 0:
        ws[f"A{NOTE}"].value = cur + text
        return
    j = len(cur)
    for nm in next_markers:
        k = cur.find(nm, i + 1)
        if k >= 0:
            j = min(j, k)
    ws[f"A{NOTE}"].value = cur[:i] + text + cur[j:]


def style(c, *, fill=None, bold=False, color=None, fmt=None, center=False, wrap=False, size=10):
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.font = Font(bold=bold, color=color or "FF000000", size=size)
    if fmt:
        c.number_format = fmt
    if center or wrap:
        c.alignment = Alignment(horizontal="center" if center else None,
                                vertical="center", wrap_text=wrap)


def headers(ws):
    return {ws.cell(5, c).value: GL(c) for c in range(1, ws.max_column + 2)
            if ws.cell(5, c).value}


def add_section(ws, census, label):
    rev = census.get("revision", "LB1")
    nruns = len(census.get("runs_used") or [])
    H = headers(ws)
    need = ["Database", "Query Qs", "All Qs", "Query P1 %", "Query P2 %",
            "Success Rate (P1)", "Blended P2 %"]
    missing = [n for n in need if n not in H]
    if missing:
        raise SystemExit(f"{ws.title}: no {missing} column - layout not recognised")
    raw = "Raw Query P1 %" in H and "Raw Query P2 %" in H

    # Rebuild the section from scratch if it is already there, so the script is
    # idempotent and a corrected census can simply be re-applied.
    start = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(4, c).value or "").startswith("CENSUS"):
            start = c
            break
    if start is None:
        last_used = max((c for c in range(1, ws.max_column + 1)
                         if any(ws.cell(r, c).value is not None for r in range(1, NOTE))),
                        default=ws.max_column)
        start = last_used + 2  # one separator column, as the other sections use
        ws.column_dimensions[GL(start - 1)].width = 2
    # Only this section's own row-4 merge. `min_col >= start` also caught the
    # section headers to the RIGHT of this one and left them unmerged, which is
    # how the corrected-regime and cost banners lost their formatting.
    end_hdr = next((c for c in range(start + 1, ws.max_column + 1)
                    if ws.cell(4, c).value), ws.max_column + 1)
    for rng in [r for r in ws.merged_cells.ranges
                if r.min_row == 4 and start <= r.min_col < end_hdr]:
        ws.unmerge_cells(str(rng))
    # Clear rows 4..ALL only. Row 29 (BEST SINGLE RUN) carries this section's
    # adjusted rates for the quoted run, written by add_best_run_breakdown.py as
    # formulas over the ALL row - wiping them here left row 29 blank every time
    # the census was re-assembled.
    end_clear = next((c for c in range(start + 1, ws.max_column + 1)
                      if ws.cell(4, c).value), ws.max_column + 2)
    for c in range(start, end_clear):
        for r in range(4, ALL + 1):
            ws.cell(r, c).value = None

    cols, c = {}, start

    def add(name, head, width, fmt=None):
        nonlocal c
        cols[name] = (GL(c), head, fmt)
        ws.column_dimensions[GL(c)].width = width
        c += 1

    add("u1", "Query P1 unwinnable Qs", 12, "0")
    add("u2", "Query P2 unwinnable Qs", 12, "0")
    add("u1p", "Query P1 unwinnable %", 12, PCT)
    add("aq1", "Query P1 Adjusted %", 12, PCT)
    add("aq2", "Query P2 Adjusted %", 12, PCT)
    add("asr", "Success Rate Adjusted (P1)", 14, PCT)
    add("ap2", "Blended P2 Adjusted", 13, PCT)
    if raw:
        add("ar1", "Raw Query P1 Adjusted %", 13, PCT)
        add("ar2", "Raw Query P2 Adjusted %", 13, PCT)
        add("lift1", "Query P1 adjusted lift", 13, "+0%;-0%")
        add("lift2", "Query P2 adjusted lift", 13, "+0%;-0%")
    add("brd", "Board-broken golds (P1)", 12, "0")
    end = c - 1

    cell = ws.cell(4, start)
    cell.value = (f"CENSUS-ADJUSTED - UNWINNABLE IN LEADERBOARD MODE "
                  f"(revision {rev}, {nruns} runs)")
    style(cell, fill=LILAC, bold=True, color=NAVY, center=True, wrap=True)
    ws.merge_cells(start_row=4, start_column=start, end_row=4, end_column=end)
    for _, (letter, head, _f) in cols.items():
        cl = ws[f"{letter}5"]
        cl.value = head
        style(cl, fill=NAVY, bold=True, color="FFFFFFFF", center=True, wrap=True)

    L = {k: v[0] for k, v in cols.items()}
    for r in range(FIRST, LAST + 1):
        db = ws[f"{H['Database']}{r}"].value
        k = census["per_db"].get(db)
        if not k:
            raise SystemExit(f"{ws.title} row {r}: {db!r} not in the census")
        u1 = k.get("q_unw_p1", 0)
        u2 = k.get("q_unw_p2", 0)
        brd = k.get("q_board_p1", 0) + k.get("m_board_p1", 0)
        write_row(ws, r, L, H, u1, u2, brd, raw)
    # pooled row: sums of the counts, and the same ratio formulas over row 28
    ws[f"{L['u1']}{ALL}"] = f"=SUM({L['u1']}{FIRST}:{L['u1']}{LAST})"
    ws[f"{L['u2']}{ALL}"] = f"=SUM({L['u2']}{FIRST}:{L['u2']}{LAST})"
    ws[f"{L['brd']}{ALL}"] = f"=SUM({L['brd']}{FIRST}:{L['brd']}{LAST})"
    write_rates(ws, ALL, L, H, raw)
    # best single run: the tab's own row-29 figures over the pooled denominators
    for key, src, cnt, den in (("asr", "Success Rate (P1)", "u1", "All Qs"),
                               ("ap2", "Blended P2 %", "u2", "All Qs")):
        ws[f"{L[key]}{BEST}"] = (
            f'=IF({H[den]}{ALL}>{L[cnt]}{ALL},{H[src]}{BEST}*{H[den]}{ALL}'
            f'/({H[den]}{ALL}-{L[cnt]}{ALL}),"n/a")')
    for r in (ALL, BEST):
        for c2 in range(start, end + 1):
            cl = ws.cell(r, c2)
            cl.fill = PatternFill("solid", fgColor=BAND)
            cl.font = Font(bold=True, size=10)
    for r in (FIRST, ALL, BEST):
        for name, (letter, _h, fmt) in cols.items():
            if fmt:
                ws[f"{letter}{r}"].number_format = fmt
    for r in range(FIRST, LAST + 1):
        for name, (letter, _h, fmt) in cols.items():
            if fmt:
                ws[f"{letter}{r}"].number_format = fmt

    upsert_note(ws, MARKER, HOW_TO_READ.format(rev=rev, runs=nruns),
                next_markers=(" LIFT IS QUERY-ONLY",))
    ws.row_dimensions[NOTE].height = 150
    return end - start + 1


def write_row(ws, r, L, H, u1, u2, brd, raw):
    ws[f"{L['u1']}{r}"] = u1
    ws[f"{L['u2']}{r}"] = u2
    ws[f"{L['brd']}{r}"] = brd
    write_rates(ws, r, L, H, raw)


def write_rates(ws, r, L, H, raw):
    q, n = H["Query Qs"], H["All Qs"]
    ws[f"{L['u1p']}{r}"] = f'=IF({q}{r}=0,"n/a",{L["u1"]}{r}/{q}{r})'
    pairs = [("aq1", "Query P1 %", "u1", q), ("aq2", "Query P2 %", "u2", q),
             ("asr", "Success Rate (P1)", "u1", n), ("ap2", "Blended P2 %", "u2", n)]
    if raw:
        pairs += [("ar1", "Raw Query P1 %", "u1", q), ("ar2", "Raw Query P2 %", "u2", q)]
    for key, src, cnt, den in pairs:
        ws[f"{L[key]}{r}"] = (
            f'=IF({den}{r}>{L[cnt]}{r},{H[src]}{r}*{den}{r}/({den}{r}-{L[cnt]}{r}),"n/a")')
    if raw:
        # The lift over the census-adjusted rates. It is ARITHMETICALLY IDENTICAL
        # to the plain Query lift: both arms divide by the same (Qs - unwinnable),
        # so the winnable-task denominator cancels in the ratio. Carried because
        # the adjusted rates are what the census section reports, and a reader
        # should be able to see the lift in the same units without taking the
        # cancellation on trust. ISNUMBER guards because the cells it divides are
        # themselves formulas that can return the text "n/a".
        for key, a_key, r_key in (("lift1", "aq1", "ar1"), ("lift2", "aq2", "ar2")):
            ws[f"{L[key]}{r}"] = (
                f'=IF(AND(ISNUMBER({L[a_key]}{r}),ISNUMBER({L[r_key]}{r}),'
                f'{L[r_key]}{r}<>0),{L[a_key]}{r}/{L[r_key]}{r}-1,"n/a")')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--census", required=True)
    ap.add_argument("--src", default=DRIVE_XLSX)
    ap.add_argument("--out")
    ap.add_argument("--in-place", action="store_true")
    a = ap.parse_args()
    if not a.out and not a.in_place:
        raise SystemExit("give --out <file.xlsx> or --in-place")
    census = json.load(open(a.census))
    wb = openpyxl.load_workbook(a.src)
    tabs = [n for n in wb.sheetnames if n.startswith("Leaderboard")]
    if not tabs:
        raise SystemExit("no Leaderboard tabs")
    for n in tabs:
        cnt = add_section(wb[n], census, n)
        print(f"{n}: {cnt} census columns")
    out = a.out or a.src
    if a.in_place:
        b = Path(out).with_name(Path(out).stem + f".bak_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        shutil.copy2(a.src, b)
        print(f"backup: {b}")
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
