#!/usr/bin/env python3
"""Add a COST (USD) section to every "Leaderboard *" tab.

What the harness actually records: results/llm_usage.jsonl carries one row per LLM
call with ts, role, model and cost_usd -- and NO task or database id. Each run file
then rolls that up into llm_usage.total / by_role. So dollar cost exists per RUN and
per ROLE, and nowhere else.

That is why this section leaves the 22 database rows blank and writes only the ALL
row. With concurrency 5 the calls of five tasks interleave in the ledger, so a
per-database figure could only be produced by bucketing calls to tasks by timestamp
and guessing which of the five in flight each belonged to. A column of numbers that
looks per-database but is really an apportionment would be read as measured, so the
cells stay empty until the harness stamps a task id on each call.

The bird-coins section beside this one IS per-database: coins are BIRD's own cost
model, recorded per task by the benchmark. Coins answer "what did the protocol
charge", dollars answer "what did it cost to run" -- they are not interchangeable,
which is the reason to show both.

The raw arm's run file is not named in the tab's provenance paragraph, so it is
matched on agent_model: the raw control for a tab is the most recent raw run whose
agent is the same model the tab's atscale runs used.

    PYTHONPATH=. .venv-adk/bin/python scripts/add_cost_columns.py --in-place
"""
from __future__ import annotations
import argparse, glob, json, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as GL

DRIVE_XLSX = ("/Users/davidmariani/Library/CloudStorage/GoogleDrive-dave@atscale.com/"
              "Shared drives/Product/Benchmark/BIRD Benchmark/Results/"
              "BIRD Results (atscale n=3, memory=shapes, mcp=feedback-memory-combined, engine=pr9967).xlsx")
RESULTS = Path("results")
FIRST, LAST, ALL, BEST, NOTE = 6, 27, 28, 29, 30
USD, USD3, NUM1 = '"$"#,##0.00', '"$"#,##0.000', "0.0"
TEAL, NAVY = "FFD6E9E4", "FF1F3864"
MARKER = " COST (added 2026-09-18"
HOW_TO_READ = (
    " COST (added 2026-09-18). Dollar cost of the run, from the harness's own LLM ledger "
    "(results/llm_usage.jsonl, one row per call). It is recorded per RUN and per ROLE and "
    "carries no task or database id, so only the ALL row can be filled: with concurrency 5 "
    "the calls of five tasks interleave, and splitting them by timestamp would be an "
    "apportionment presented as a measurement. The bird-coins section is the per-database "
    "cost view -- coins are what the BIRD protocol charges, dollars are what the run cost to "
    "execute. USD per P1 correct divides the mean run cost by the tasks that run got right, "
    "so it prices an answer rather than an attempt.")


def headers(ws):
    return {ws.cell(5, c).value: GL(c) for c in range(1, ws.max_column + 2)
            if ws.cell(5, c).value}


def tab_runs(ws):
    return re.findall(r"[\w.\-]+\.json", ws["A2"].value or "")


def usage(path):
    d = json.load(open(path))
    u = (d.get("llm_usage") or {}).get("total") or {}
    R = d.get("results") or []
    scored = [t for t in R if "phase1_passed" in t]
    return {
        "cost": u.get("cost_usd"), "calls": u.get("calls"),
        "tasks": len(R) or None,
        "p1": sum(1 for t in scored if t["phase1_passed"]),
        "agent": d.get("agent_model"),
    }


def raw_for(agent):
    """The raw control for a tab: newest raw run with the same agent model."""
    best = None
    for p in sorted(glob.glob("results/leaderboard_*_raw_*.json")):
        try:
            d = json.load(open(p))
        except Exception:
            continue
        if d.get("agent_model") == agent:
            best = p
    return usage(best) if best else None


def upsert_note(ws, marker, text):
    cur = ws[f"A{NOTE}"].value or ""
    if marker in cur:
        return
    ws[f"A{NOTE}"].value = cur + text


def add_section(ws):
    H = headers(ws)
    files = [f for f in tab_runs(ws) if (RESULTS / f).exists()]
    if not files:
        print(f"{ws.title}: no run files found - skipped")
        return 0
    U = [usage(RESULTS / f) for f in files]
    U = [u for u in U if u["cost"]]
    if not U:
        print(f"{ws.title}: run files carry no llm_usage cost - skipped")
        return 0
    raw = raw_for(U[0]["agent"])

    # idempotent: rebuild in place if the section is already there
    start = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(4, c).value or "").startswith("COST"):
            start = c
            break
    if start is None:
        start = ws.max_column + 2
        ws.column_dimensions[GL(start - 1)].width = 2
    for rng in [r for r in ws.merged_cells.ranges if r.min_row == 4 and r.min_col >= start]:
        ws.unmerge_cells(str(rng))
    for c in range(start, ws.max_column + 2):
        for r in range(4, NOTE):
            ws.cell(r, c).value = None

    cols, c = {}, start

    def add(name, width, fmt):
        nonlocal c
        cell = ws.cell(5, name and c)
        cell.value = name
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[GL(c)].width = width
        cols[name] = (GL(c), fmt)
        c += 1

    for i in range(1, len(U) + 1):
        add(f"USD r{i}", 11, USD)
    add("USD mean/run", 13, USD)
    add("USD per task", 12, USD3)
    add("USD per P1 correct", 15, USD3)
    if raw:
        add("Raw USD", 11, USD)
        add("Raw USD per task", 14, USD3)

    head = ws.cell(4, start)
    head.value = "COST (USD) - run-level; the ledger carries no per-database id"
    head.fill = PatternFill("solid", fgColor=TEAL)
    head.font = Font(bold=True, color=NAVY, size=10)
    head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if c - 1 > start:
        ws.merge_cells(start_row=4, start_column=start, end_row=4, end_column=c - 1)

    mean = sum(u["cost"] for u in U) / len(U)
    mean_tasks = sum(u["tasks"] for u in U) / len(U)
    mean_p1 = sum(u["p1"] for u in U) / len(U)
    vals = {f"USD r{i}": u["cost"] for i, u in enumerate(U, 1)}
    vals["USD mean/run"] = mean
    vals["USD per task"] = mean / mean_tasks if mean_tasks else None
    vals["USD per P1 correct"] = mean / mean_p1 if mean_p1 else None
    if raw:
        vals["Raw USD"] = raw["cost"]
        vals["Raw USD per task"] = raw["cost"] / raw["tasks"] if raw["tasks"] else None

    for name, (letter, fmt) in cols.items():
        cell = ws[f"{letter}{ALL}"]
        cell.value = vals.get(name)
        cell.number_format = fmt
        cell.font = Font(bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor="FFFFF2CC")

    upsert_note(ws, MARKER, HOW_TO_READ)
    return len(cols)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=DRIVE_XLSX)
    ap.add_argument("--out")
    ap.add_argument("--in-place", action="store_true")
    a = ap.parse_args()
    if not a.out and not a.in_place:
        raise SystemExit("give --out <file.xlsx> or --in-place")
    wb = openpyxl.load_workbook(a.src)
    for n in [s for s in wb.sheetnames if s.startswith("Leaderboard")]:
        cnt = add_section(wb[n])
        if cnt:
            print(f"{n}: {cnt} cost columns")
    out = a.src if a.in_place else a.out
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
