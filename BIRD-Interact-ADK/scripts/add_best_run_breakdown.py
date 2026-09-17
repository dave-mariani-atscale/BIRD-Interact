#!/usr/bin/env python3
"""Fill the BEST SINGLE RUN row of every "Leaderboard *" tab with the Query and
Management breakdown, not just the blended figures.

Row 29 is the leaderboard reporting convention - one run, the best of the three
by blended reward - but the tab only carried its blended Success Rate, Blended
P2, Reward and Efficiency. The per-category columns (Query P1/P2 %, Mgmt P1/P2 %)
existed only as the task-mean on row 28, so the split behind the headline number
could not be read for the run actually being quoted. There are no per-run
category columns on the tab to INDEX into, so the four rates are written as
literals from that run's own results file.

The run is chosen exactly as row 29 already chooses it: the highest pooled
blended reward (0.7*P1 + 0.3*P2 over all 600 tasks). Every section of row 29
therefore describes the SAME run - blended, per-category, census-adjusted and,
where a regrade file is supplied, the corrected-regime comparison.

    PYTHONPATH=. .venv-adk/bin/python scripts/add_best_run_breakdown.py \
        --regrade '../../bird-atscale-models/census/regrade_09*.json' --in-place
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import json
import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as GL

DRIVE_XLSX = ("/Users/davidmariani/Library/CloudStorage/GoogleDrive-dave@atscale.com/"
              "Shared drives/Product/Benchmark/BIRD Benchmark/Results/"
              "BIRD Results 22db — 2026-09-01 (atscale n=3, Semantic Memory shapes, engine pr9967).xlsx")
REV5 = ("/Users/davidmariani/workspace/atscale/bird-atscale-models/census/"
        "census_rev5_final.json")
RESULTS = Path("results")
FIRST, LAST, ALL, BEST, NOTE = 6, 27, 28, 29, 30
PCT, NUM = "0.0%", "0.0"
MARKER = " ROW 29 CARRIES THE CATEGORY SPLIT"
NOTE_ADD = (
    " ROW 29 CARRIES THE CATEGORY SPLIT (added 2026-09-17). The best-single-run row now also shows "
    "that run's Query P1/P2 % and Mgmt P1/P2 %, its census-adjusted Query rates, and - where the "
    "corrected-regime section is present - the same run on those rules. The run is the one row 29 "
    "already quotes: the highest pooled blended reward of the three, named in the Coins used cell. "
    "The four category rates are literals read from that run's results file, because the tab has no "
    "per-run category columns to INDEX into; everything derived from them is still a formula.")


def tab_runs(ws):
    return [f for f in re.findall(r"[\w.\-]+\.json", ws["A2"].value or "")]


def run_stats(path):
    d = json.load(open(path))
    q = m = None
    agg = {"q": [0, 0, 0], "m": [0, 0, 0]}   # n, p1, p2
    for r in d["results"]:
        cat = r.get("category") or ("Management" if "_M_" in r["task_id"] else "Query")
        k = "q" if cat == "Query" else "m"
        agg[k][0] += 1
        agg[k][1] += 1 if r.get("phase1_passed") else 0
        agg[k][2] += 1 if r.get("phase2_passed") else 0
    per_task = {(r.get("instance_id") or r["task_id"]):
                (1 if r.get("phase1_passed") else 0, 1 if r.get("phase2_passed") else 0)
                for r in d["results"]}
    del d
    n = agg["q"][0] + agg["m"][0]
    p1 = agg["q"][1] + agg["m"][1]
    p2 = agg["q"][2] + agg["m"][2]
    return {
        "n": n, "reward": (0.7 * p1 + 0.3 * p2) / n if n else 0,
        "sr": p1 / n if n else 0, "b_p2": p2 / n if n else 0,
        "q_p1": agg["q"][1] / agg["q"][0], "q_p2": agg["q"][2] / agg["q"][0],
        "m_p1": agg["m"][1] / agg["m"][0], "m_p2": agg["m"][2] / agg["m"][0],
        "q_n": agg["q"][0], "per_task": per_task,
    }


def census():
    r5 = json.load(open(REV5))
    unw1 = {t for t, v in r5.items() if v["tier"] in ("T1", "T2")}
    unw2 = {t for t, v in r5.items()
            if v["tier"] in ("T1", "T2") or v.get("gold_p2") or v.get("content_p2")
            or v.get("order_p2")}
    return r5, unw1, unw2


def headers(ws):
    return {ws.cell(5, c).value: GL(c) for c in range(1, ws.max_column + 2)
            if ws.cell(5, c).value}


def upsert_note(ws, marker, text):
    cur = ws[f"A{NOTE}"].value or ""
    if marker in cur:
        return
    ws[f"A{NOTE}"].value = cur + text


def put(ws, letter, row, value, fmt):
    cell = ws[f"{letter}{row}"]
    cell.value = value
    cell.number_format = fmt
    cell.font = Font(bold=True, size=10)


def fill_best(ws, regrades):
    H = headers(ws)
    files = tab_runs(ws)
    stats = {f: run_stats(RESULTS / f) for f in files if (RESULTS / f).exists()}
    if not stats:
        print(f"{ws.title}: no results files found - skipped")
        return 0
    best = max(stats, key=lambda f: stats[f]["reward"])
    idx = files.index(best) + 1
    s = stats[best]

    for key, col in (("q_p1", "Query P1 %"), ("q_p2", "Query P2 %"),
                     ("m_p1", "Mgmt P1 %"), ("m_p2", "Mgmt P2 %")):
        if col in H:
            put(ws, H[col], BEST, s[key], PCT)
    written = 4

    # census-adjusted Query rates for the same run: formulas off the cells above
    q, b = H.get("Query Qs"), H.get("Query P1 %")
    for adj, src, unw in (("Query P1 Adjusted %", "Query P1 %", "Query P1 unwinnable Qs"),
                          ("Query P2 Adjusted %", "Query P2 %", "Query P2 unwinnable Qs")):
        if adj in H and unw in H and q:
            ws[f"{H[adj]}{BEST}"] = (
                f'=IF(SUM({q}{FIRST}:{q}{LAST})>{H[unw]}{ALL},'
                f'{H[src]}{BEST}*SUM({q}{FIRST}:{q}{LAST})/'
                f'(SUM({q}{FIRST}:{q}{LAST})-{H[unw]}{ALL}),"n/a")')
            ws[f"{H[adj]}{BEST}"].number_format = PCT
            ws[f"{H[adj]}{BEST}"].font = Font(bold=True, size=10)
            written += 1

    # corrected-regime section, same run, if the regrade data covers it
    rg = next((r for r in regrades if best in r.get("per_task", {})), None)
    if rg and "Query P1 % (corrected)" in H:
        _r5, unw1, unw2 = census()
        pt = rg["per_task"][best]
        tot = {"p1": [0, 0], "p2": [0, 0]}  # all, achievable-only
        for tid in _r5:
            for ph, unw in (("p1", unw1), ("p2", unw2)):
                cell = pt.get(f"{tid}|{ph}")
                v = cell["corrected"] if cell else s["per_task"].get(tid, (0, 0))[0 if ph == "p1" else 1]
                tot[ph][0] += v
                if tid not in unw:
                    tot[ph][1] += v
        put(ws, H["Query P1 % (corrected)"], BEST, tot["p1"][0] / len(_r5), PCT)
        put(ws, H["Query P2 % (corrected, floor)"], BEST, tot["p2"][0] / len(_r5), PCT)
        put(ws, H["P1 Qs passed, achievable only"], BEST, tot["p1"][1], NUM)
        put(ws, H["P2 Qs passed, achievable only"], BEST, tot["p2"][1], NUM)
        for adj, ach, unw in (("Query P1 Adjusted % (corrected)",
                               "P1 Qs passed, achievable only", "P1 unsolvable Qs (rev 5)"),
                              ("Query P2 Adjusted % (corrected, floor)",
                               "P2 Qs passed, achievable only", "P2 unsolvable Qs (rev 5)")):
            ws[f"{H[adj]}{BEST}"] = (
                f'=IF(SUM({q}{FIRST}:{q}{LAST})>{H[unw]}{ALL},'
                f'{H[ach]}{BEST}/(SUM({q}{FIRST}:{q}{LAST})-{H[unw]}{ALL}),"n/a")')
            ws[f"{H[adj]}{BEST}"].number_format = PCT
            ws[f"{H[adj]}{BEST}"].font = Font(bold=True, size=10)
        written += 6

    # the builder's own label cell names the run; make it explicit rather than a formula
    if "Coins used" in H:
        cell = ws[f"{H['Coins used']}{BEST}"]
        cell.value = f"best run by reward: r{idx} ({best[:46]})"
        cell.font = Font(bold=True, size=9)

    upsert_note(ws, MARKER, NOTE_ADD)
    print(f"{ws.title}: best = r{idx} {best}  "
          f"Query P1 {s['q_p1']:.1%} / P2 {s['q_p2']:.1%}, "
          f"Mgmt P1 {s['m_p1']:.1%} / P2 {s['m_p2']:.1%}, "
          f"blended SR {s['sr']:.2%} reward {s['reward']:.4f}")
    return written


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--regrade", nargs="*", default=[])
    ap.add_argument("--src", default=DRIVE_XLSX)
    ap.add_argument("--out")
    ap.add_argument("--in-place", action="store_true")
    a = ap.parse_args()
    if not a.out and not a.in_place:
        raise SystemExit("give --out <file.xlsx> or --in-place")
    regrades = [json.load(open(p)) for g in a.regrade for p in sorted(glob.glob(g))]
    wb = openpyxl.load_workbook(a.src)
    for n in [s for s in wb.sheetnames if s.startswith("Leaderboard")]:
        cells = fill_best(wb[n], regrades)
        print(f"  -> {cells} cells on row {BEST}")
    out = a.out or a.src
    if a.in_place:
        b = Path(out).with_name(Path(out).stem + f".bak_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        shutil.copy2(a.src, b)
        print(f"backup: {b}")
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
