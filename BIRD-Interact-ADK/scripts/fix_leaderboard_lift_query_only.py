#!/usr/bin/env python3
"""Retarget the LIFT section of every "Leaderboard *" tab onto QUERY tasks only.

The tabs were built with a blended lift: Success Rate lift (all 600 tasks) and
Reward lift (all 600). That understates the thing it is trying to measure. In
leaderboard mode the 190 Management tasks are routed to the raw Postgres tools
on BOTH arms - the same system twice, by construction - so a third of the
denominator can only ever pull the ratio toward 1.0. The semantic layer answers
Query tasks and nothing else, so the lift belongs there.

After this the section reads:

  Query P1 lift       Query P1 % / Raw Query P1 % - 1
  Query P2 lift       Query P2 % / Raw Query P2 % - 1
  Query reward lift   (0.7*P1 + 0.3*P2) over the Query columns, both arms

Census note: a lift is a RATIO of two rates over the same task set, so the
census-adjusted rates give exactly the same lift - the (Qs - unwinnable) factor
cancels top and bottom. The census block carries that pair anyway ("Query P1 /
P2 adjusted lift", scripts/add_leaderboard_census_columns.py) so the lift can be
read in the same units as the adjusted rates; the two sections agreeing to the
last digit is the check that the cancellation is real.

Idempotent: it rewrites the three lift columns and their header in place, so
re-running it after a rebuild is safe.

    PYTHONPATH=. .venv-adk/bin/python scripts/fix_leaderboard_lift_query_only.py --in-place
"""
from __future__ import annotations

import argparse
import datetime as dt
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as GL

DRIVE_XLSX = ("/Users/davidmariani/Library/CloudStorage/GoogleDrive-dave@atscale.com/"
              "Shared drives/Product/Benchmark/BIRD Benchmark/Results/"
              "BIRD Results 22db — 2026-09-01 (atscale n=3, Semantic Memory shapes, engine pr9967).xlsx")
FIRST, LAST, ALL, NOTE = 6, 27, 28, 30
NAVY, BLUE = "FF1F3864", "FFBDD7EE"
FMT = "+0%;-0%"
HEADS = ["Query P1 lift", "Query P2 lift", "Query reward lift"]
SECTION = "LIFT - ATSCALE vs RAW (QUERY TASKS ONLY)"
NOTE_ADD = (
    " LIFT IS QUERY-ONLY (changed 2026-09-16). Management tasks run on the raw Postgres tools on "
    "both arms in leaderboard mode, so a blended lift compares the semantic layer against itself on "
    "a third of the benchmark; these three columns divide the AtScale Query rate by the raw Query "
    "rate, and Query reward lift rebuilds BIRD's own 0.7*P1 + 0.3*P2 from those two columns. A lift "
    "is a ratio of two rates over the same tasks, so the census-adjusted rates give the identical "
    "lift - the winnable-task denominator cancels - and the census block's Query P1/P2 adjusted lift "
    "columns should read the same to the last digit.")


def upsert_note(ws, marker, text, next_markers=()):
    """Replace this script's own paragraph in the HOW TO READ cell, leaving the
    census script's paragraph alone - the two run in either order."""
    cur = ws[f"A{NOTE}"].value or ""
    i = cur.find(marker)
    if i < 0:
        ws[f"A{NOTE}"].value = cur + text
        return
    j = len(cur)
    for nm in next_markers:
        k = cur.find(nm, i + 1)
        if k >= 0:
            j = min(j, k)
    ws[f"A{NOTE}"].value = cur[:i] + text + cur[j:]


def headers(ws):
    return {ws.cell(5, c).value: GL(c) for c in range(1, ws.max_column + 1)
            if ws.cell(5, c).value}


def fix(ws):
    H = headers(ws)
    need = ["Query P1 %", "Query P2 %", "Raw Query P1 %", "Raw Query P2 %"]
    missing = [n for n in need if n not in H]
    if missing:
        print(f"{ws.title}: skipped, no {missing}")
        return 0
    # The lift block is wherever the P1 lift column already sits; the two columns
    # after it are the ones being retargeted (Success Rate / Reward lift, or this
    # script's own output on a re-run).
    start_letter = H.get("Query P1 lift")
    if not start_letter:
        print(f"{ws.title}: skipped, no 'Query P1 lift' column")
        return 0
    start = ws[f"{start_letter}5"].column
    cols = [GL(start + i) for i in range(3)]

    cell = ws.cell(4, start)
    cell.value = SECTION
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(bold=True, color=NAVY, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for letter, head in zip(cols, HEADS):
        c = ws[f"{letter}5"]
        c.value = head
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color="FFFFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    a1, a2, r1, r2 = (H["Query P1 %"], H["Query P2 %"],
                      H["Raw Query P1 %"], H["Raw Query P2 %"])
    for r in list(range(FIRST, LAST + 1)) + [ALL]:
        pairs = ((cols[0], a1, r1), (cols[1], a2, r2))
        for letter, a, raw in pairs:
            c = ws[f"{letter}{r}"]
            c.value = f'=IF(OR({raw}{r}="",{raw}{r}=0),"n/a",{a}{r}/{raw}{r}-1)'
            c.number_format = FMT
        num = f"(0.7*{a1}{r}+0.3*{a2}{r})"
        den = f"(0.7*{r1}{r}+0.3*{r2}{r})"
        c = ws[f"{cols[2]}{r}"]
        c.value = f'=IF(OR({r1}{r}="",{den}=0),"n/a",{num}/{den}-1)'
        c.number_format = FMT

    upsert_note(ws, " LIFT IS QUERY-ONLY", NOTE_ADD,
                next_markers=(" CENSUS-ADJUSTED (added",))
    return 3


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=DRIVE_XLSX)
    ap.add_argument("--out")
    ap.add_argument("--in-place", action="store_true")
    a = ap.parse_args()
    if not a.out and not a.in_place:
        raise SystemExit("give --out <file.xlsx> or --in-place")
    wb = openpyxl.load_workbook(a.src)
    for n in [s for s in wb.sheetnames if s.startswith("Leaderboard")]:
        print(f"{n}: {fix(wb[n])} lift columns retargeted")
    out = a.out or a.src
    if a.in_place:
        b = Path(out).with_name(Path(out).stem + f".bak_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        shutil.copy2(a.src, b)
        print(f"backup: {b}")
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
